"""Build the defense deck skeleton (Deck v2 in docs/defense/outline.md) as a PPTX.

Layout mirrors the peer FPT deck (docs/defense/reference/LVTN.pptx): I-IV sections, sub-dividers in III,
hierarchical numbering, footer bar with section name + page number, THE END, backups.
Colours are ours (#1F2933 bars), not theirs. Coordinates are written in px on a
1920x1080 canvas (= Canva's canvas); font sizes are in pt on the 1440x810 pt page, which
is exactly how Canva exports, so a size here equals the size Canva shows after import.

Speaker notes carry the talk-path marker: "★ TALK m:ss" or "○ SUPPORT", plus the thesis source.

Usage: .venv/Scripts/python.exe docs/defense/scripts/build_deck.py
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUT = ROOT / "deck_v2.pptx"

FONT = "Segoe UI"
DARK = RGBColor(0x1F, 0x29, 0x33)     # bars + main text
MUTED = RGBColor(0x6B, 0x72, 0x80)    # placeholder labels
PH_FILL = RGBColor(0xF7, 0xF7, 0xF7)
PH_LINE = RGBColor(0xD1, 0xD5, 0xDB)
CARD_FILL = RGBColor(0xEE, 0xF0, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

X0, X1 = 120, 1800          # content left/right
Y0, Y1 = 220, 960           # content top/bottom
GAP = 40


def px(v):
    return Emu(int(round(v * 9525)))


# ---------------------------------------------------------------- primitives
def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, lines, size, bold=False, italic=False, color=DARK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, upper=False, space_after=0):
    """lines: str or list of str / (str, dict-overrides)."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if isinstance(lines, str):
        lines = [lines]
    for i, item in enumerate(lines):
        s, ov = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(ov.get("space_after", space_after))
        r = p.add_run()
        r.text = s.upper() if upper else s
        f = r.font
        f.name = FONT
        f.size = Pt(ov.get("size", size))
        f.bold = ov.get("bold", bold)
        f.italic = ov.get("italic", italic)
        f.color.rgb = ov.get("color", color)
    return tb


def placeholder(slide, x, y, w, h, label):
    rect(slide, x, y, w, h, PH_FILL, PH_LINE)
    text(slide, x + 20, y, w - 40, h, f"[ {label} ]", 20, italic=True, color=MUTED,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(slide, x, y, w, h, head, body):
    rect(slide, x, y, w, h, CARD_FILL)
    text(slide, x + 30, y + 40, w - 60, h - 80,
         [(head, {"bold": True, "size": 32, "space_after": 14}), (body, {"size": 28})],
         28, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def image_fit(slide, path, x, y, w, h):
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    dw, dh = iw * s, ih * s
    slide.shapes.add_picture(str(path), px(x + (w - dw) / 2), px(y + (h - dh) / 2), px(dw), px(dh))
    return x + (w - dw) / 2, y + (h - dh) / 2, dw, dh


def footer(slide, section, number):
    rect(slide, 0, 1000, 1920, 80, DARK)
    if section:
        text(slide, 40, 1000, 1200, 80, section, 24, bold=True, italic=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE, upper=True)
    text(slide, 1600, 1000, 280, 80, str(number), 24, color=WHITE,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)


def top_bar(slide):
    rect(slide, 1240, 50, 600, 45, DARK)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


# ---------------------------------------------------------------- layouts
def split(n, x=X0, w=X1 - X0):
    cw = (w - GAP * (n - 1)) / n
    return [(x + i * (cw + GAP), cw) for i in range(n)]


def lay_full(sl, labels):
    placeholder(sl, X0, Y0, X1 - X0, Y1 - Y0, labels[0])


def lay_two(sl, labels, ratio=0.5):
    w = X1 - X0 - GAP
    wl = w * ratio
    placeholder(sl, X0, Y0, wl, Y1 - Y0, labels[0])
    placeholder(sl, X0 + wl + GAP, Y0, w - wl, Y1 - Y0, labels[1])


def lay_cards_ph(sl, labels):
    for (x, w), lab in zip(split(len(labels)), labels):
        placeholder(sl, x, Y0, w, Y1 - Y0, lab)


def lay_cards_text(sl, items):
    for (x, w), (h, b) in zip(split(len(items)), items):
        card(sl, x, Y0, w, Y1 - Y0, h, b)


def lay_columns(sl, cols, item_size=32):
    for (x, w), (head, items) in zip(split(len(cols)), cols):
        text(sl, x, Y0, w, 60, head, 36, bold=True)
        text(sl, x, Y0 + 80, w, Y1 - Y0 - 80, items, item_size, space_after=18)


def lay_image(sl, name):
    image_fit(sl, ASSETS / name, X0, Y0, X1 - X0, Y1 - Y0)


def lay_image_or_ph(sl, arg):
    """(asset name, placeholder label): the user-drawn export once it exists, else a placeholder."""
    name, label = arg
    if (ASSETS / name).exists():
        lay_image(sl, name)
    else:
        lay_full(sl, [label])


def lay_rows_image_or_ph(sl, items):
    """Stacked full-width rows; each row is the user's draw.io export once it exists, else a
    placeholder. items: [(asset name, placeholder label), ...]"""
    h = (Y1 - Y0 - GAP * (len(items) - 1)) / len(items)
    for i, (name, label) in enumerate(items):
        y = Y0 + i * (h + GAP)
        if (ASSETS / name).exists():
            image_fit(sl, ASSETS / name, X0, y, X1 - X0, h)
        else:
            placeholder(sl, X0, y, X1 - X0, h, label)


def table(slide, x, y, w, rows, col_w, size=22, row_h=64, bold=(), align_left=1):
    """Native table: dark header row, light zebra body. rows[0] = header; bold = {(r, c)};
    the first `align_left` columns are left-aligned, the rest centred."""
    shp = slide.shapes.add_table(len(rows), len(rows[0]), px(x), px(y), px(w), px(row_h * len(rows)))
    tbl = shp.table
    for c, cw in enumerate(col_w):
        tbl.columns[c].width = px(w * cw / sum(col_w))
    for r, row in enumerate(rows):
        tbl.rows[r].height = px(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = DARK if r == 0 else (CARD_FILL if r % 2 == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = px(12)
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c < align_left else PP_ALIGN.CENTER
            run = p.add_run()
            run.text = val
            f = run.font
            f.name, f.size = FONT, Pt(size)
            f.bold = r == 0 or (r, c) in bold
            f.color.rgb = WHITE if r == 0 else DARK
    return y + row_h * len(rows)


def lay_clustering_tables(sl, _):
    """3.6.2: tab:clustering (Mean IoU column left out) + tab:splitrate, numbers verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.56
    xr = X0 + wl + GAP
    text(sl, X0, Y0, wl, 50, "Clustering method (0.10 m grid, per frame, track filter off)", 24,
         bold=True)
    rows = [["Seq", "Method", "Precision", "Recall", "F1", "Time (ms)"],
            ["00", "HDBSCAN", "0.972", "0.731", "0.835", "306"],
            ["00", "DBSCAN", "0.975", "0.637", "0.771", "54"],
            ["00", "Euclidean", "0.979", "0.636", "0.771", "76"],
            ["08", "HDBSCAN", "0.840", "0.671", "0.746", "320"],
            ["08", "DBSCAN", "0.870", "0.565", "0.685", "57"],
            ["08", "Euclidean", "0.869", "0.576", "0.693", "75"]]
    yb = table(sl, X0, Y0 + 70, wl, rows, [0.7, 1.5, 1.3, 1.1, 1.0, 1.3], align_left=2,
               bold={(1, 1), (1, 3), (1, 4), (4, 1), (4, 3), (4, 4)})
    text(sl, X0, yb + 24, wl, 120,
         "Precision within 0.03; HDBSCAN gains 9–11 recall points for about 6× DBSCAN's clustering time.", 24)
    text(sl, xr, Y0, w - wl, 50, "Clustering grid (seq 08, 34,057 cars)", 24, bold=True)
    rows = [["", "0.05 m", "0.10 m"],
            ["Car in one cluster", "55.1%", "66.0%"],
            ["Car split", "44.8%", "33.7%"],
            ["Car merged", "0.2% (56)", "0.8% (284)"],
            ["Clusters per split car", "4.7", "3.0"]]
    yb = table(sl, xr, Y0 + 70, w - wl, rows, [1.9, 1.1, 1.2], bold={(2, 2), (3, 2)})
    text(sl, xr, yb + 24, w - wl, 160,
         ["Recall 0.710 → 0.732 (+756 true positives),", "F1 0.798 → 0.812."], 24)


def lay_results(sl, _):
    """3.3.1: D9 stage chart (tab:ablation-stage) + headline table (tab:headline, tab:loso), verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.58
    image_fit(sl, ASSETS / "d9_stage_funnel.png", X0, Y0, wl, Y1 - Y0)
    xr = X0 + wl + GAP
    text(sl, xr, Y0 + 40, w - wl, 50, "Detection results", 24, bold=True)
    rows = [["", "Precision", "Recall", "F1"],
            ["Seq 08", "0.911", "0.732", "0.812"],
            ["All 11 sequences", "0.888", "0.717", "0.793"]]
    yb = table(sl, xr, Y0 + 110, w - wl, rows, [2.4, 1.2, 1.0, 0.9], size=21, row_h=78,
               bold={(1, 0), (1, 1), (1, 2), (1, 3)})
    text(sl, xr, yb + 24, w - wl, 200,
         ["All 11 sequences are cross-validated: each is scored by a classifier trained without it.",
          "Within about 2 points of seq 08 on every metric."], 21, space_after=10)


def lay_geo_filter(sl, _):
    """2.2.3: thesis fig:geo-filter crop + all gates (ch3 §Geometric Filtering list), verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.54
    image_fit(sl, ASSETS / "thesis_geo_filter.png", X0, Y0 + 60, wl, Y1 - Y0 - 120)
    xr = X0 + wl + GAP
    rows = [["Gate", "Keep the cluster if"],
            ["Points", "≥ 15"],
            ["Box volume", "0.5–50 m³"],
            ["Longest side", "≤ 6.0 m"],
            ["Size floors", "longest ≥ 0.5 m, middle ≥ 0.4 m"],
            ["Centre height", "≤ 1.5 m above the ground plane"],
            ["Vertical span", "≤ 1.8 m"],
            ["Aspect ratio", "longest / shortest ≤ 6.0"]]
    yb = table(sl, xr, Y0, w - wl, rows, [1.0, 1.9], size=21, row_h=62, align_left=2)
    text(sl, xr, yb + 20, w - wl, 120,
         ["After this filter alone: precision 0.199, recall 0.773.",
          "Loose by design: the classifier makes the fine decision."], 21, space_after=8)


def lay_stage_ablation(sl, _):
    """3.6.1: tab:ablation-stage, verbatim (full seq 08, frozen configuration)."""
    rows = [["Configuration", "Precision", "Recall", "F1", "Mean IoU", "TP / FP / FN"],
            ["Geometric only", "0.199", "0.773", "0.316", "0.907", "27,002 / 108,831 / 7,920"],
            ["+ classifier", "0.846", "0.676", "0.751", "0.922", "23,600 / 4,303 / 11,322"],
            ["+ classifier, track filter", "0.911", "0.732", "0.812", "0.913", "25,568 / 2,513 / 9,354"]]
    yb = table(sl, X0, Y0 + 40, X1 - X0, rows, [2.3, 1.1, 1.0, 0.9, 1.1, 2.6], size=24, row_h=80,
               bold={(3, 0), (3, 1), (3, 2), (3, 3)})
    text(sl, X0, yb + 30, X1 - X0, 200,
         ["The classifier is the precision step: false positives 108,831 → 4,303, at a recall cost.",
          "The track filter wins recall back: true positives 23,600 → 25,568.",
          "First two rows are per frame; the last row uses the offline track filter. "
          "Mean IoU is over matched detections only."], 24, space_after=10)


def lay_pretraining(sl, _):
    """3.6.4: ch4 sec:ablation-stagea numbers + tab:cross-domain, verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.5
    xr = X0 + wl + GAP
    text(sl, X0, Y0, wl, 50, "Classifier level (seq 08 validation clusters)", 24, bold=True)
    rows = [["Training recipe", "Macro-F1"],
            ["Real clusters only, from scratch (final)", "0.9285"],
            ["Synthetic pretraining, then real", "0.9225"]]
    yb = table(sl, X0, Y0 + 70, wl, rows, [2.6, 1.0], size=22, row_h=72, bold={(1, 0), (1, 1)})
    text(sl, X0, yb + 24, wl, 160,
         ["Within run-to-run noise. At pipeline level the two give identical F1 to three decimals."], 22)
    text(sl, xr, Y0, w - wl, 50, "Cross-domain transfer (car-class F1)", 24, bold=True)
    rows = [["Trained on ↓ / tested on →", "Synthetic", "Real"],
            ["Synthetic (Stage A)", "0.999", "0.000"],
            ["Real (Stage B)", "0.000", "0.885"]]
    yb = table(sl, xr, Y0 + 70, w - wl, rows, [2.2, 1.1, 1.0], size=22, row_h=72,
               bold={(1, 2), (2, 1)})
    text(sl, xr, yb + 24, w - wl, 160,
         ["A model from one domain finds no cars in the other: the gap is total and symmetric."], 22)


def lay_literature(sl, _):
    """1.3: three lines of prior work, wording from ch2 §2.2 (bg-deep, bg-classical, bg-completion)."""
    rows = [["Approach", "Representative works", "Strengths", "Remaining challenges"],
            ["End-to-end deep segmentation", "RangeNet++, RandLA-Net, Cylinder3D; panoptic: Panoptic-PolarNet, DS-Net",
             "Benchmark leaders (Cylinder3D 67.8 mIoU on the SemanticKITTI test set)",
             "Need dense per-point (and per-instance) labels; a missed object can't be traced to one step"],
            ["Classical geometric segmentation", "RANSAC ground removal + DBSCAN / HDBSCAN clustering",
             "No training data for segmentation; every stage inspectable",
             "No notion of objectness: clustering splits cars, which costs recall"],
            ["Object-level completion", "PCN, PoinTr, SnowflakeNet, SeedFormer",
             "Learns to map a partial view to a complete shape",
             "Trained on synthetic CAD shapes (domain gap); Chamfer scoring needs a complete true shape that real LiDAR lacks"]]
    table(sl, X0, Y0 + 20, X1 - X0, rows, [1.25, 1.75, 1.6, 2.1], size=21, row_h=125, align_left=4,
          bold={(1, 0), (2, 0), (3, 0)})


def lay_datasets(sl, _):
    """3.1.1: tab:dataset-roles, tab:sensor-spec, tab:dataset, verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.6
    xr = X0 + wl + GAP
    rows = [["", "SemanticKITTI", "ShapeNetCore v2"],
            ["Data", "Real LiDAR scans", "Synthetic CAD meshes"],
            ["Role", "Detection, mining, evaluation", "Completion network training"],
            ["Volume", "23,201 frames (11 sequences)", "~3,150 car models"],
            ["Annotations used", "Point-wise semantic + instance", "None (synthetic partials)"]]
    table(sl, X0, Y0, wl, rows, [1.1, 1.6, 1.5], size=20, row_h=58)
    rows = [["Sensor", "Velodyne HDL-64E"],
            ["Beams", "64"],
            ["Frame rate", "10 Hz"],
            ["Horizontal resolution", "~0.09°"],
            ["Max range", "120 m"],
            ["Points per frame", "~100,000–120,000"]]
    table(sl, xr, Y0, w - wl, rows, [1.4, 1.3], size=20, row_h=48)
    seqs = ["00", "01", "02", "03", "04", "05", "06", "07", "08", "09", "10"]
    frames = ["4541", "1101", "4661", "801", "271", "2761", "1101", "1101", "4071", "1591", "1201"]
    rows = [["Seq"] + seqs, ["Frames"] + frames, ["Classifier"] + ["val" if q == "08" else "train" for q in seqs]]
    yb = table(sl, X0, Y0 + 370, X1 - X0, rows, [1.6] + [1.0] * 11, size=20, row_h=56,
               bold={(1, 9), (2, 9)})
    text(sl, X0, yb + 18, X1 - X0, 100,
         ["Sequences 11–21 have no public labels, so no labelled sequence is left as a separate test set. "
          "For completion, every sequence is evaluation only: PCN trains on synthetic ShapeNet."], 20)


def lay_problem(sl, _):
    w = X1 - X0 - GAP
    wl = w * 0.6
    image_fit(sl, ASSETS / "d1_occlusion.png", X0, Y0, wl, Y1 - Y0)
    xr = X0 + wl + GAP
    ix, iy, iw, ih = image_fit(sl, ASSETS / "scan_panel_b.png", xr, Y0, w - wl, 680)
    text(sl, xr, iy + ih + 12, w - wl, 40, "Real car, 6 m away (Thesis Fig. 2.3)", 24,
         align=PP_ALIGN.CENTER)


def lay_eligibility(sl, _):
    """3.1.2: ch4 §Target Eligibility rule and stride-20 counts, verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.45
    xr = X0 + wl + GAP
    text(sl, X0, Y0 + 20, wl, Y1 - Y0 - 40,
         ["Ground truth: SemanticKITTI's own car and moving-car instances.",
          "A car is scored only if at least 10 of its points survive preprocessing and ground removal.",
          "With fewer points no clustering method can segment it: counting it would measure sensor "
          "sparsity, not the pipeline."], 26, space_after=24)
    rows = [["Seq 08 sample (204 frames)", "Car instances"],
            ["Annotated", "2,332"],
            ["Eligible (≥ 10 surviving points)", "1,737"],
            ["Excluded", "25.5%"]]
    yb = table(sl, xr, Y0 + 20, w - wl, rows, [2.2, 1.0], size=22, row_h=66, bold={(3, 1)})
    rows = [["Recall measured against", "Recall"],
            ["Eligible cars (headline)", "0.732"],
            ["All annotated cars (inferred)", "≈ 0.55"]]
    yb = table(sl, xr, yb + 40, w - wl, rows, [2.2, 1.0], size=22, row_h=66, bold={(1, 1)})
    text(sl, xr, yb + 16, w - wl, 60, "0.732 × (1,737 / 2,332) ≈ 0.55", 22, color=MUTED)


def lay_metrics(sl, _):
    """3.2.2: thesis fig:bg-iou + ch4 §Evaluation Metrics; completion metric = ch4 §plausibility."""
    w = X1 - X0 - GAP
    wl = w * 0.58
    xr = X0 + wl + GAP
    ix, iy, iw, ih = image_fit(sl, ASSETS / "thesis_bg_iou.png", X0, Y0, wl, 520)
    text(sl, X0, Y0 + 540, wl, Y1 - Y0 - 540,
         ["Best overlap first, one-to-one: matched = TP, unmatched detection = FP, missed car = FN.",
          "Counts are summed over every frame of the sequence."], 22, space_after=8)
    card(sl, xr, Y0, w - wl, Y1 - Y0, "Completion: plausibility",
         "An independent judge, a PointNet++ classifier trained on ModelNet40 shapes, labels each "
         "completed car. Score = the share it calls a car.")


def lay_loso(sl, _):
    """3.2.3: LOSO grid (scripts/make_loso_grid.py) + tab:headline / tab:loso rows, verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.5
    xr = X0 + wl + GAP
    image_fit(sl, ASSETS / "loso_grid.png", X0, Y0, wl, Y1 - Y0)
    rows = [["", "Precision", "Recall", "F1"],
            ["Seq 08, final classifier", "0.911", "0.732", "0.812"],
            ["Seq 08, blind fold", "0.888", "0.736", "0.805"],
            ["All 11 folds, pooled", "0.888", "0.717", "0.793"]]
    yb = table(sl, xr, Y0 + 40, w - wl, rows, [2.3, 1.2, 1.0, 0.9], size=22, row_h=72,
               bold={(3, 0), (3, 1), (3, 2), (3, 3)})
    text(sl, xr, yb + 30, w - wl, 260,
         ["Each fold: 15 epochs, last checkpoint kept, no selection on the test sequence.",
          "Choosing the final checkpoint on seq 08 bought about 2 precision points and no recall."],
         22, space_after=12)


def lay_recall_limit(sl, _):
    """3.3.3: D10 (thesis fig:failure-zooms, frame-250 row) + split rate and recall by range
    (ch4 §split analysis at 0.10 m; tab:distance recall), verbatim."""
    ix, iy, iw, ih = image_fit(sl, ASSETS / "d10_split_car_frame250.png", X0, Y0, X1 - X0, 500)
    rows = [["Range", "0–10 m", "10–20 m", "20–40 m", "40 m+"],
            ["Cars split by clustering", "53.5%", "40.2%", "22.3%", "3.5%"],
            ["Recall (per frame)", "0.787", "0.862", "0.550", "0.196"]]
    table(sl, X0, iy + ih + 24, X1 - X0, rows, [2.4, 1, 1, 1, 1], size=22, row_h=60,
          bold={(1, 1), (2, 2)})


def lay_distance(sl, _):
    """3.3.4: tab:distance, verbatim (stride-20, per frame, track filter off)."""
    rows = [["Range (m)", "Precision", "Recall", "F1", "Mean IoU", "TP / FP / FN"],
            ["0–10", "0.899", "0.787", "0.839", "0.902", "311 / 35 / 84"],
            ["10–20", "0.869", "0.862", "0.866", "0.926", "432 / 65 / 69"],
            ["20–40", "0.799", "0.550", "0.652", "0.934", "401 / 101 / 328"],
            ["40+", "0.512", "0.196", "0.284", "0.943", "22 / 21 / 90"],
            ["Pooled", "0.840", "0.671", "0.746", "0.923", "1,166 / 222 / 571"]]
    yb = table(sl, X0, Y0 + 20, X1 - X0, rows, [1.2, 1.1, 1.0, 0.9, 1.1, 2.0], size=24, row_h=70,
               bold={(2, 2)})
    text(sl, X0, yb + 26, X1 - X0, 200,
         ["Nearest cars: lower recall than 10–20 m despite more points, because half of them are split.",
          "Beyond 20 m: ordinary sparsity. Per frame with the track filter off, so read it as a trend."],
         24, space_after=10)


def lay_runtime(sl, _):
    """3.5.1: tab:runtime, verbatim."""
    w = X1 - X0 - GAP
    wl = w * 0.6
    xr = X0 + wl + GAP
    rows = [["Stage", "Mean (ms)", "Share"],
            ["Data load", "0.7", "0.1%"],
            ["Preprocessing", "125.9", "20.2%"],
            ["Ground removal (RANSAC)", "56.2", "9.0%"],
            ["Clustering (HDBSCAN)", "370.4", "59.4%"],
            ["Geometric filter", "36.2", "5.8%"],
            ["Classification", "33.7", "5.4%"],
            ["Tracking", "0.3", "0.1%"],
            ["Total", "623.5", "≈ 1.6 fps"]]
    table(sl, X0, Y0, wl, rows, [2.4, 1.1, 1.1], size=22, row_h=78,
          bold={(4, 0), (4, 1), (4, 2), (8, 0), (8, 1), (8, 2)})
    text(sl, xr, Y0 + 40, w - wl, Y1 - Y0 - 80,
         ["Full seq 08, 4,068 timed frames.",
          "Classical CPU stages: about 94% of the time.",
          "The classifier, the only learned detection stage: 5.4%.",
          "Offline by design: real time is out of scope."], 24, space_after=20)


def lay_runtime_opt(sl, _):
    """3.5.2: ch3 §RANSAC (1000 -> 300) + ch4 §runtime evaluation, verbatim."""
    rows = [["", "Before", "After"],
            ["RANSAC iterations", "1,000", "300"],
            ["Clustering grid", "0.05 m", "0.10 m"],
            ["Time per scan (full seq 08)", "921.3 ms", "623.5 ms"]]
    yb = table(sl, X0 + 200, Y0 + 20, X1 - X0 - 400, rows, [2.2, 1.2, 1.2], size=26, row_h=84,
               bold={(3, 2)})
    text(sl, X0 + 200, yb + 30, X1 - X0 - 400, 240,
         ["About a third faster; the coarser grid also raised recall (0.710 → 0.732, slide 45).",
          "Below 400 ms was not reachable: Python HDBSCAN alone has a floor of about 185 ms."],
         24, space_after=12)


def lay_iou_sensitivity(sl, _):
    """3.6.3: tab:iou-sensitivity, verbatim."""
    rows = [["IoU threshold", "Precision", "Recall", "F1", "Mean IoU", "TP / FP / FN"],
            ["0.25", "0.914", "0.735", "0.815", "0.910", "25,660 / 2,421 / 9,262"],
            ["0.30 (headline)", "0.911", "0.732", "0.812", "0.913", "25,568 / 2,513 / 9,354"],
            ["0.50", "0.884", "0.711", "0.788", "0.927", "24,837 / 3,244 / 10,085"]]
    yb = table(sl, X0, Y0 + 40, X1 - X0, rows, [1.6, 1.1, 1.0, 0.9, 1.1, 2.6], size=24, row_h=80,
               bold={(2, 0), (2, 1), (2, 2), (2, 3)})
    text(sl, X0, yb + 30, X1 - X0, 200,
         ["F1 drops only 2.4 points from 0.30 to the strict 0.50: matched detections already overlap "
          "at 0.913 on average.",
          "Full seq 08, everything else frozen. Mean IoU is over matched detections only."],
         24, space_after=10)


def lay_lshape(sl, _):
    """2.3.3: thesis fig:lshape-gate + tab:completion-config values (verbatim), grouped by the D15 steps."""
    w = X1 - X0 - GAP
    wl = w * 0.40
    xr = X0 + wl + GAP
    ix, iy, iw, ih = image_fit(sl, ASSETS / "thesis_lshape_gate.png", X0, Y0 + 40, wl, 300)
    text(sl, X0, iy + ih + 40, wl, 240,
         ["Fragments and merged cars never complete into a plausible car, so they are skipped, "
          "not completed."], 22)
    rows = [["Step", "Constant", "Value", "What it corrects"],
            ["1 Check", "min. length", "2.7 m", "shorter: a fragment of a car"],
            ["", "max. width", "2.3 m", "wider: two cars merged"],
            ["2 Turn", "heading", "L-shape fit", "no constant: the fitted angle"],
            ["3 Shift", "down-shift", "0.25 m", "ground cut lifts the middle"],
            ["", "width prior", "1.9 m", "the hidden far side"],
            ["", "car length", "track p90 + 0.12 m", "the hidden far end"],
            ["", "length fallback", "4.14 m", "a track with < 5 good frames"],
            ["4 Size", "radius correction", "÷ 1.137", "one-sided radius vs whole car"]]
    table(sl, xr, Y0, w - wl, rows, [0.72, 1.2, 1.55, 2.0], size=19, row_h=74, align_left=4,
          bold={(1, 0), (3, 0), (4, 0), (8, 0)})


def mini_card(sl, x, y, w, h, head, body, size=22):
    rect(sl, x, y, w, h, CARD_FILL)
    text(sl, x + 24, y + 18, w - 48, h - 36,
         [(head, {"bold": True, "size": size + 2, "space_after": 8}), (body, {"size": size})],
         size, anchor=MSO_ANCHOR.MIDDLE)


def lay_baselines(sl, _):
    """3.2.1: ch4 §Baseline Methods + tab:positioning, as one comparison table."""
    rows = [["Part", "Compared against", "Why"],
            ["Detection stages", "each other: geometric filter → + classifier → + track filter",
             "shows what each stage adds (slides 29, 44)"],
            ["Detection vs supervised networks", "not compared head to head",
             "they score all classes (mIoU, PQ) with dense per-point labels; we score cars by point F1"],
            ["Completion", "the raw scan it started from", "the paired baseline completion must beat"],
            ["Completion", "dense multi-frame scans of parked cars", "the shape-quality upper bound"],
            ["Completion network", "PoinTr", "a heavier transformer alternative (slide 48)"]]
    table(sl, X0, Y0 + 20, X1 - X0, rows, [1.5, 2.1, 2.6], size=22, row_h=104, align_left=3,
          bold={(r, 0) for r in range(1, 6)})


def lay_judge(sl, _):
    """3.4.2: D12a (user draw.io) + D12b acceptance bars, tab:plausibility verbatim."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    w = X1 - X0 - GAP
    wl = w * 0.5
    xr = X0 + wl + GAP
    if (ASSETS / "d12a_judge.drawio.png").exists():
        image_fit(sl, ASSETS / "d12a_judge.drawio.png", X0, Y0, wl, Y1 - Y0)
    else:
        placeholder(sl, X0, Y0, wl, Y1 - Y0, "D12a judge schematic (user, draw.io)")
    cd = CategoryChartData()
    cd.categories = ["Raw scan", "Completed", "Dense reference"]
    cd.add_series("Judged a car", (0.253, 0.785, 0.807))
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, px(xr), px(Y0), px(w - wl), px(560), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Judged a car (acceptance)"
    tp = ch.chart_title.text_frame.paragraphs[0].runs[0].font
    tp.name, tp.size, tp.bold, tp.color.rgb = FONT, Pt(22), True, DARK
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 1, 0.25
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = PH_LINE
    va.tick_labels.font.size, va.tick_labels.font.name = Pt(18), FONT
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.00", False
    ca = ch.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(20), FONT
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0.000", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold, dl.font.name = Pt(22), True, FONT
    for pt, rgb in zip(plot.series[0].points, ("2B5C8A", "D9822B", "9AA0A6")):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(rgb)
    text(sl, xr, Y0 + 590, w - wl, 150,
         ["2,822 real cars, all 11 sequences; dense reference: 140 parked cars.",
          "1,615 cars rejected as a raw scan are accepted once completed; 115 go the other way."],
         21, space_after=10)


def lay_synthetic(sl, _):
    """3.4.3: D16 (scripts/make_d16_synthetic_example.py) + ch4 §synthetic fidelity numbers, verbatim."""
    ix, iy, iw, ih = image_fit(sl, ASSETS / "d16_synthetic_example.png", X0, Y0, X1 - X0, 560)
    rows = [["Held-out synthetic cars", "Chamfer distance ↓", "F-score @ 0.1 m ↑"],
            ["PCN", "0.16 m", "0.76"]]
    yb = table(sl, X0 + 260, iy + ih + 24, X1 - X0 - 520, rows, [2.0, 1.4, 1.4], size=22, row_h=58,
               bold={(1, 1), (1, 2)})
    text(sl, X0 + 260, yb + 16, X1 - X0 - 520, 50,
         "The network itself works; real scans have no true model, so they are judged separately (slide 36).",
         20, italic=True, color=MUTED)


def lay_plaus_robust(sl, _):
    """3.4.4: tab:plausibility-seq verbatim + ch4 §plausibility robustness checks, verbatim."""
    rows = [["Seq", "00", "01", "02", "03", "04", "05", "06", "07", "08", "09", "10"],
            ["Cars", "926", "96", "329", "48", "59", "258", "164", "213", "518", "141", "70"],
            ["Raw scan", "0.30", "0.07", "0.20", "0.13", "0.24", "0.33", "0.11", "0.26", "0.31", "0.13", "0.03"],
            ["Completed", "0.78", "0.89", "0.82", "0.65", "0.86", "0.74", "0.86", "0.72", "0.80", "0.70", "0.83"]]
    yb = table(sl, X0, Y0, X1 - X0, rows, [1.7] + [1] * 11, size=21, row_h=66,
               bold={(3, c) for c in range(12)})
    text(sl, X0, yb + 14, X1 - X0, 40,
         "Acceptance by sequence. The completion constants were set on seq 08; the gain holds on all eleven.",
         20, italic=True, color=MUTED)
    cards = [("Paired test", "1,615 cars flip to accepted, 115 the other way (McNemar p ≈ 0)"),
             ("Moving cars", "accepted 0.837 (602 tracks): as plausible as parked ones"),
             ("Front and back swapped", "acceptance changes by only 0.009"),
             ("Not just density", "parked cars: car probability 0.730 vs dense reference 0.728")]
    ch = min(Y1 - (yb + 80), 220)
    for (x, w), (h, b) in zip(split(4), cards):
        mini_card(sl, x, (yb + 80 + Y1 - ch) / 2, w, ch, h, b, size=22)


def lay_completion_analysis(sl, _):
    """3.4.5: tab:fidelity verbatim + ch4 §Scope of the Completion Claim."""
    w = X1 - X0 - GAP
    wl = w * 0.5
    xr = X0 + wl + GAP
    text(sl, X0, Y0, wl, 50, "Distance to the scan itself (seq 08, 39 cars)", 24, bold=True)
    rows = [["Cloud", "UCD (m) ↓", "UHD max (m) ↓"],
            ["Raw scan", "0.000", "0.000"],
            ["Mirrored scan", "0.000", "0.000"],
            ["Completed", "0.090", "0.286"]]
    yb = table(sl, X0, Y0 + 70, wl, rows, [1.6, 1.1, 1.3], size=22, row_h=72)
    text(sl, X0, yb + 24, wl, 260,
         ["Any cloud that contains the scan scores a perfect 0, so these metrics cannot credit the "
          "hidden side.",
          "Hence the independent judge (slide 36)."], 22, space_after=12)
    hc = (Y1 - Y0 - 2 * 30) / 3
    cards = [("Plausible", "Acceptance 0.253 → 0.785, on every sequence, moving and parked cars."),
             ("At the ceiling", "Accepted almost as often as dense multi-frame scans of parked cars: 0.785 vs 0.807."),
             ("Scoped claim", "Plausible car shapes on real LiDAR; not exact geometry, which no real "
                              "ground truth can score.")]
    for i, (h, b) in enumerate(cards):
        mini_card(sl, xr, Y0 + i * (hc + 30), w - wl, hc, h, b, size=22)


def lay_strategies(sl, _):
    """Backup: tab:strategies verbatim + the two prose-only strategies (ch4 §Limits of the Repair Strategies)."""
    rows = [["Strategy", "Seq 00 P", "Seq 00 R", "Seq 00 F1", "Seq 08 F1"],
            ["Baseline (HDBSCAN, min. cluster size 10)", "0.984", "0.739", "0.844", "0.829"],
            ["Bird's-eye-view clustering (0.30 m)", "0.993", "0.640", "0.779", "–"],
            ["Higher min. cluster size (20)", "0.976", "0.755", "0.852", "0.801"],
            ["Fragment merge (1.5 m, 30 points)", "0.931", "0.755", "0.834", "0.815"],
            ["Distance-adaptive HDBSCAN (30 / 15 / 10)", "0.935", "0.748", "0.831", "–"]]
    yb = table(sl, X0, Y0, X1 - X0, rows, [3.2, 1, 1, 1, 1], size=22, row_h=66,
               bold={(1, 3), (1, 4), (3, 3), (3, 4)})
    text(sl, X0, yb + 28, X1 - X0, Y1 - yb - 28,
         ["Higher min. cluster size wins on seq 00 but loses on seq 08: it erases sparse far cars (over-fitting).",
          "Fragment merge gains recall but also fuses walls and poles: precision 0.984 → 0.931.",
          "Also tried, in the text only: stacking 3 frames before clustering (F1 0.073); lower size thresholds (no new TP).",
          "What worked instead: a coarser clustering grid, 0.05 → 0.10 m (slide 45)."], 22, space_after=10)


def lay_judge_independent(sl, _):
    """Backup: ch4 §Real Data Plausibility + findings #53/#54, as threat -> control."""
    rows = [["Could the judge favour us?", "What rules it out"],
            ["It learned from the same shapes as PCN", "Trained on ModelNet40, a different repository: it never saw a ShapeNet model"],
            ["We graded our own work", "Not the pipeline's classifier: a separate PointNet++, trained from scratch"],
            ["The judge is weak", "Test accuracy 0.9214, above a pass mark of 0.88 fixed in advance"],
            ["It calls everything a car", "It rejects 3 in 4 raw scans (acceptance 0.253)"],
            ["It just rewards dense clouds", "Dense real cars score 0.807; parked completions match them (car prob. 0.730 vs 0.728)"],
            ["The reading was chosen after the fact", "Hypotheses and the meaning of every outcome were written before scoring"]]
    table(sl, X0, Y0, X1 - X0, rows, [1.45, 2.6], size=22, row_h=100, align_left=2,
          bold={(r, 0) for r in range(1, 7)})


def lay_ucd_uhd(sl, _):
    """Backup: toy figure (scripts/make_b59_ucd_uhd.py) + ch2 sec:th-distances, tab:fidelity."""
    ix, iy, iw, ih = image_fit(sl, ASSETS / "b59_ucd_uhd.png", X0, Y0, X1 - X0, 560)
    text(sl, X0 + 80, iy + ih + 30, X1 - X0 - 160, Y1 - (iy + ih + 30),
         ["For each scan point, the distance to the nearest output point: UCD averages them, UHD takes the largest.",
          "Real seq-08 cars: the raw scan scores 0.000, the completion UCD 0.090 m (slide 39). "
          "Chen et al. pair UHD with a plausibility judge for this reason."], 22, space_after=10)


def lay_iou_example(sl, _):
    """Backup: ch2 eq. IoU + ch4 §Evaluation Metrics; worked example = S4 notes A4 [illustration]."""
    w = X1 - X0 - GAP
    wl = w * 0.36
    xr = X0 + wl + GAP
    text(sl, X0, Y0, wl, Y1 - Y0,
         [("IoU = shared points ÷ all points of the two", {"bold": True}),
          "1. Every detection–car pair with IoU ≥ 0.3 is a candidate.",
          "2. Best IoU first; each detection and each car is used once.",
          "3. Matched = true positive; extra detection = false positive; missed car = false negative.",
          "4. Counted over every frame, then precision, recall, F1."], 22, space_after=16)
    text(sl, xr, Y0, w - wl, 50, "One car of 400 points (illustration)", 24, bold=True)
    rows = [["Case", "IoU", "Outcome"],
            ["One cluster of 360 points, 340 on the car", "340 / (400 + 360 − 340) = 0.81", "1 TP"],
            ["Split into 250 + 150 points", "0.625 and 0.375", "1 TP + 1 FP"],
            ["Split into four of 100 points", "0.25 each, below 0.3", "1 FN (+ FPs)"]]
    yb = table(sl, xr, Y0 + 70, w - wl, rows, [2.1, 1.9, 1.1], size=20, row_h=84, align_left=1)
    text(sl, xr, yb + 24, w - wl, 200,
         ["Seq 08: 25,568 TP, 2,513 FP, 9,354 FN → precision 0.911, recall 0.732, F1 0.812."], 22)


def lay_pipeline_config(sl, _):
    """Backup: tab:pipeline-config, verbatim, in two halves."""
    w = X1 - X0 - GAP
    half = w / 2
    left = [["Stage", "Parameter", "Value"],
            ["Preprocess", "z-filter threshold", "−2.0 m"],
            ["", "voxel size", "0.05 m"],
            ["", "denoise neighbours / std ratio", "20 / 2.0"],
            ["", "voxel before denoise", "true"],
            ["Ground", "RANSAC iterations", "300"],
            ["", "inlier distance", "0.2 m"],
            ["", "min. normal z", "0.5"],
            ["Clustering", "method", "HDBSCAN"],
            ["", "min. cluster size / min. samples", "10 / 5"],
            ["", "clustering voxel size", "0.10 m"]]
    right = [["Stage", "Parameter", "Value"],
             ["Geometric", "min. points", "15"],
             ["", "volume range", "[0.5, 50] m³"],
             ["", "max. longest extent", "6.0 m"],
             ["", "min. top two extents", "0.5 m / 0.4 m"],
             ["", "max. height above ground", "1.5 m"],
             ["", "max. vertical span", "1.8 m"],
             ["", "max. aspect ratio", "6.0"],
             ["Classifier", "model", "dual-branch PointNet"],
             ["", "input points", "512"],
             ["Tracker", "max. match distance", "2.0 m"],
             ["", "max. disappeared frames", "5"],
             ["", "min. track length", "2"],
             ["", "min. known votes / ratio", "2 / 0.5"]]
    rh = (Y1 - Y0) / len(right)
    for x, rows, cw in ((X0, left, [1.1, 2.5, 1.2]), (X0 + half + GAP, right, [1.1, 2.0, 1.7])):
        table(sl, x, Y0, half, rows, cw, size=18, row_h=rh, align_left=2,
              bold={(r, 0) for r in range(1, len(rows))})


def lay_completion_ablations(sl, _):
    """3.6.5: ch4 §Completion Ablation Studies, one row per ablation."""
    rows = [["Choice", "Tried instead", "Result"],
            ["Car-frame normalisation", "3D PCA frame, own centre and radius",
             "same network makes mis-oriented, mis-scaled blobs"],
            ["KITTI-like training scans", "fine-tune on noisier simulated scans",
             "did not recover car shape: dropped (negative result)"],
            ["L-shape gate", "complete every car",
             "on seq 08, every implausible completion came from a fragment or merge"],
            ["Per-car length", "one fixed 4.14 m length",
             "stretches compact cars, cuts long ones short"],
            ["PCN", "PoinTr (transformer)", "Chamfer 0.161 → 0.153 m only: PCN kept"]]
    table(sl, X0, Y0 + 20, X1 - X0, rows, [1.5, 1.9, 2.6], size=22, row_h=104, align_left=3,
          bold={(r, 0) for r in range(1, 6)})


# ---------------------------------------------------------------- slide builders
def content(prs, n, section, title, key, layout, arg, note):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    text(sl, X0, 30, X1 - X0, 70, title, 44, bold=True, align=PP_ALIGN.CENTER, upper=True)
    if key:
        text(sl, X0, 105, X1 - X0, 100, key, 30, align=PP_ALIGN.CENTER)
    layout(sl, arg)
    footer(sl, section, n)
    notes(sl, note)


def divider(prs, n, line1, line2=None, note="divider: click through", sub=False):
    """Section divider (line1 at 56) or, with sub=True, a mid-section sub-divider whose only
    line uses the same style as a section divider's second line (44)."""
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    top_bar(sl)
    lines = [(line1, {"size": 44 if sub else 56})] + ([(line2, {"size": 44})] if line2 else [])
    text(sl, X0, 300, X1 - X0, 400, lines, 56, bold=True, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE, upper=True, space_after=12)
    footer(sl, None, n)
    notes(sl, note)


def cover(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    image_fit(sl, ASSETS / "fpt_logo.png", 120, 40, 300, 110)
    top_bar(sl)
    text(sl, 160, 280, 1600, 330,
         "A Modular Detection and Completion Framework for Occluded Vehicles in Automotive LiDAR",
         48, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, upper=True)
    text(sl, 1100, 720, 760, 220,
         [("Supervisor: Dr. Doan Nhat Quang", {"italic": True}),
          "Presenter: Ngo Vi Viet Anh", "ID: 24MSE13205"],
         30, bold=True, space_after=6)
    footer(sl, None, 1)
    notes(sl, "★ TALK 0:15 — title, name, supervisor.")


def toc(prs):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    top_bar(sl)
    text(sl, X0, 150, X1 - X0, 90, "Table of Contents", 56, bold=True,
         align=PP_ALIGN.CENTER, upper=True)
    text(sl, 300, 400, 640, 300, ["I. Introduction", "II. Proposed Methodology"], 36,
         bold=True, space_after=60)
    text(sl, 1040, 400, 700, 300, ["III. Experimental Evaluation", "IV. Conclusion & Future Work"],
         36, bold=True, space_after=60)
    footer(sl, None, 2)
    notes(sl, "★ TALK 0:10 — four parts.")


def end(prs, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    image_fit(sl, ASSETS / "fpt_logo.png", 810, 60, 300, 110)
    text(sl, X0, 200, X1 - X0, 140,
         [("Thank you.", {"size": 56}), ("Questions are welcome.", {"size": 36})],
         56, bold=True, align=PP_ALIGN.CENTER, upper=True, space_after=8)
    image_fit(sl, ASSETS / "d3_master_pipeline.drawio.png", X0, 400, X1 - X0, 560)  # D3: Q&A anchor
    footer(sl, None, n)
    notes(sl, "★ END — leave up during Q&A; jump to slides by number.")


# ---------------------------------------------------------------- deck content
I, II, III, IV, BK = ("I. Introduction", "II. Proposed Methodology",
                      "III. Experimental Evaluation", "IV. Conclusion & Future Work", "Backup")
T = lambda t, src: f"★ TALK {t} — source: {src}"     # noqa: E731
S = lambda src: f"○ SUPPORT (skip in talk) — source: {src}"  # noqa: E731


def add_vi_script(prs):
    """Append each ★ slide's VI line from talk_script.md to its speaker notes (presenter view).
    The script's '### <n> <title> (m:ss)' heading must match the deck: same number, same title."""
    import re
    blocks, cur = {}, None
    for line in (ROOT / "talk_script.md").read_text(encoding="utf-8").splitlines():
        m = re.match(r"### (\d+) (.+) \((\d+:\d\d)\)$", line)
        if m:
            cur = int(m[1])
            blocks[cur] = {"title": m[2]}
        elif cur and line.startswith("VI: "):
            blocks[cur]["vi"] = line[4:]
    for n, b in blocks.items():
        sl = prs.slides[n - 1]
        if n not in (1, 53):
            title = next(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip())
            assert title.upper() == b["title"].upper(), f"talk_script.md slide {n}: '{b['title']}' vs deck '{title}'"
        tf = sl.notes_slide.notes_text_frame
        tf.text = tf.text + "\n\nVI: " + b["vi"]


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    nxt = lambda: len(prs.slides) + 1  # noqa: E731  (auto page numbers)
    C = lambda sec, title, key, lay, arg, note: content(prs, nxt(), sec, title, key, lay, arg, note)  # noqa: E731
    D = lambda l1, l2=None: divider(prs, nxt(), l1, l2)  # noqa: E731
    SUB = lambda l1: divider(prs, nxt(), l1, sub=True)  # noqa: E731

    cover(prs)
    toc(prs)

    # ---------------- I. INTRODUCTION (friend: 1 Background ... 5 Objectives & Contributions)
    D("I. Introduction")
    C(I, "1.1 Background", "The laser only reaches the near side, so half of every car is missing.",
      lay_problem, None, T("1:00", "D1; fig:bg-lidar-scan (b); say the 3 point-cloud properties aloud (backup slide)"))
    C(I, "1.2 Current Approaches",
      "End-to-end networks are accurate but need dense labels and act as black boxes; we take the modular route.",
      lay_image, "d2_end_to_end_vs_modular.png",
      T("1:00", "D2; ch1 §1.1 (examples); ch2 §2.2.2–2.2.3. If asked how we compare: different metrics "
               "(their all-class mIoU vs our car-only point F1), not one scale; see 1.3 / tab:positioning"))
    C(I, "1.3 Literature Review", "Three lines of prior work, and what each leaves open.",
      lay_literature, None,
      S("ch2 §2.2 (bg-deep, bg-classical, bg-completion); tab:positioning for the per-method table. "
        "No accuracy column on purpose: their mIoU / PQ over all classes and our car-only point F1 are "
        "different measurements (ch2 §positioning)."))
    C(I, "1.4 Research Gap",
      "Accurate detectors need dense labels and are hard to inspect, and real-world completion has no ground truth to train or test against.",
      lay_cards_text, [("Dense labels", "Accurate networks need a class per point or a box per object"),
                       ("Opacity", "A missed car can't be traced to one step"),
                       ("No paired real data", "No real occluded car comes with its full shape"),
                       ("No ground truth", "Real completions can't be scored against a true shape")],
      T("0:40", "ch1 §1.1–1.2"))
    C(I, "1.5 Objectives & Contributions", "Three objectives, two contributions.",
      lay_columns, [("Objectives", [
          "1. Test whether a modular, geometry-first pipeline can detect vehicles as a transparent alternative to end-to-end networks.",
          "2. Evaluate point cloud completion on real LiDAR, where no complete ground truth exists.",
          "3. Test whether synthetic-only training and auto-mined supervision transfer to real LiDAR without new manual labels."]),
          ("Contributions", [
          "1. A modular detection and completion framework for occluded vehicles, inspectable and reproducible at every stage.",
          "2. Zero additional manual annotation: classifier labels are mined automatically; the completion network trains only on synthetic shapes."])],
      T("0:45", "ch1 §Objectives, §Contributions"))

    # ---------------- II. PROPOSED METHODOLOGY (friend: 1 -> 2 -> 2.1..2.4 -> 3 -> 4 -> 5)
    D("II. Proposed Methodology")
    C(II, "2.1 System Architecture", "Eight steps from a raw scan to completed cars.",
      lay_image, "d3_master_pipeline.drawio.png", T("1:00", "D3; fig:architecture, fig:pipeline, fig:completion-pipeline"))
    C(II, "2.2 Detection Pipeline", "Six steps turn a raw scan into tracked cars.",
      lay_image, "thesis_pipeline.png",
      S("ch3 §detection pipeline; fig:pipeline (crop: scripts/crop_thesis_figure.py 3.2 pipeline)"))
    C(II, "2.2.1 Preprocessing and Ground Removal",
      "Clean the scan, then remove the ground plane with RANSAC.",
      lay_image, "d4_ransac_real.png",
      T("0:45", "D4 (seq 08 frame 2500, one-frame illustration: counts are NOT thesis results); ch2 §RANSAC; ch3 §preprocessing, §RANSAC."))
    C(II, "2.2.2 Clustering (HDBSCAN)", "HDBSCAN groups the remaining points into object candidates.",
      lay_image, "d5_ground_then_clusters.png",
      T("0:45", "D5 = same seq 08 frame 2500 as D4 (scripts/make_d5_clusters.py, pipeline's own "
               "preprocess, remove_ground, cluster_objects); one-frame illustration: the 205 clusters are NOT "
               "a thesis result. Concept: fig:clustering; fig:bg-pipeline-stages."))
    C(II, "2.2.3 Geometric Filter",
      "Size and shape rules on each cluster's oriented box drop obvious non-cars.",
      lay_geo_filter, None,
      S("fig:geo-filter (crop: scripts/crop_thesis_figure.py 3.4 geo_filter); gates = ch3 §Geometric Filtering; "
        "P 0.199 / R 0.773 = tab:ablation-stage geometric-only row. If asked which gates cost cars: 27,718 rejected "
        "clusters overlap a GT car, 73.9% by the minimum-volume and 12.7% by the minimum-points gate, i.e. small "
        "splinters of already-split cars (the thesis states this as an inference)."))
    C(II, "2.2.4 Car / Not-Car Classifier",
      "A small PointNet decides car or not-car, trained on labels mined automatically from SemanticKITTI.",
      lay_image_or_ph, ("d67_classifier_training.drawio.png",
                        "D6+D7 merged: training lane trains the classifier (draw.io: assets/d67_classifier_training.drawio)"),
      T("1:00", "D6+D7 merged. Top lane (training) = ch3 §training data (mined, not hand labelled); bottom lane "
               "(inference) = fig:classifier, simplified. Story: the top lane is the answer key, looked up from "
               "SemanticKITTI labels; the classifier learns to give the same answer from the points alone, so on "
               "a new scan, with no labels, only the bottom lane runs. "
               "Say 'zero additional manual annotation', not 'annotation-free'. "
               "Counts: 420,333 training clusters (seq 00-07, 09, 10), 130,394 validation (seq 08), "
               "about 21% car / 79% not-car."))
    C(II, "2.2.5 Tracking and Track Filter",
      "Detections are linked across frames; each track is kept or dropped as a whole, by vote.",
      lay_image_or_ph, ("d8_tracking_timeline.drawio.png",
                        "D8 tracking timeline (scaffold: assets/d8_tracking_timeline.drawio)"),
      T("0:30", "D8; fig:tracker, ch3 §Multi-Frame Tracking, tab:pipeline-config tracker rows. Say: each "
                "detection is linked to the nearest track within 2 metres, frame to frame. Then one decision "
                "per track, not per frame: a track is a car if at least 2 and at least half of its frames vote "
                "car. Top: a parked car the classifier misses once is kept in every frame, so recall goes up. "
                "Bottom: a bush called a car once is dropped, so precision goes up. On seq 08: +1,968 true "
                "positives and 1,790 fewer false positives (tab:ablation-stage). It uses the whole track, so "
                "it is offline, like the rest of the system."))
    C(II, "2.3 Completion Pipeline",
      "The network trains only on synthetic shapes but must work on real scans.",
      lay_image_or_ph, ("d11_pcn.drawio.png", "D11 PCN training and inference (scaffold: assets/d11_pcn.drawio)"),
      T("1:00", "D11; fig:pcn, ch3 §Completion Network, §Training Data, §Inference-Time Partial Normalisation. Say: PCN is "
                "a standard encoder-decoder: it squeezes the partial into one shape code and grows a full car from it, 256 "
                "points in, 4,096 out. It is trained only on synthetic ShapeNet meshes, so two things make it work on real "
                "scans. First, the training inputs are simulated to look like a real scan: one laser viewpoint, 5 cm voxels, "
                "the ground cut off (next slide). Second, at test time the real scan is checked, turned, shifted and sized "
                "into exactly the frame the network trained in (slide 19). At first the network produced only blobs. "
                "The fix was not more training. (Stop there.) HOOK H1 → backup 57; answer in qa_bank.md."))
    C(II, "2.3.1 Synthetic Training Partials",
      "Training inputs are simulated to look like real scans.",
      lay_image, "d14_training_vs_real.png", S("D14 (scripts/make_d14_partials.py): illustration, not a result. "
      "Panel 1 = the training code's own generator on the thesis fig:kitti-like-partial mesh (one HDL-64E view, 5 cm "
      "voxels, 0.30 m ground cut); panel 2 = the complete-model sample PCN must output (ShapeNet meshes include the "
      "interior, hence the seats); panel 3 = a real single-frame scan (car C, slide 35). Say: the simulated input has "
      "the same one-sided view and scan lines as the real scan, so the network meets the same kind of input on the "
      "road as in training. The network sees 256 of these points. ch3 §Training Data. If asked why not a simpler "
      "render: an early dense pinhole render had no scan lines, no voxels and no ground cut, so it did not look like "
      "LiDAR. Training: 100 epochs, Adam 1e-4, batch 8, best validation loss 0.1246."))
    C(II, "2.3.2 Real Scan Normalisation",
      "Before PCN, each real car is checked, turned, shifted and sized.",
      lay_image, "d15_check_turn_shift_size.png",
      S("D15 (scripts/make_d15_normalisation.py): the production geometry code on car C (seq 03, slides 18 and "
        "35), the exact scan production completed. Illustration, not a result. 1 Check: L-shape fit 3.7 x 1.5 m, "
        "inside the gate (at least 2.7 m long, at most 2.3 m wide). 2 Turn: length onto the training axis, gravity "
        "up. 3 Shift: the scan's middle is not the car's middle; push 0.95 - 1.52/2 = 0.19 m away from the sensor "
        "(1.9 m width prior), 0.23 m toward the unseen end, 0.25 m down for the ground cut. 4 Size: divide by the "
        "largest distance / 1.137; points can poke past the circle, that's the correction. Then 256 points go into "
        "PCN and the result is rotated back. Memory hook: check, turn, shift, size. ch3 §Object Completion, "
        "§Inference-Time Partial Normalisation, tab:completion-config."))
    C(II, "2.3.3 Gate and Normalisation Constants",
      "How the gate decides, and what each constant corrects.",
      lay_lshape, None, S("Close-up of slide 19's four steps. fig:lshape-gate (crop_thesis_figure.py 3.10 lshape_gate); "
        "values = tab:completion-config verbatim, grouped by step. Where the numbers come from: 1.137 is the median "
        "ratio of a one-sided partial's radius to the true radius, measured on synthetic cars; 4.14 m is the median "
        "car length; 1.9 m a full car width; 0.25 m undoes the upward bias of the ground cut (ch3). "
        "L-shape fit (Zhang 2017): try rectangle orientations 1 degree apart and keep the one whose edges the points "
        "hug; longer side = length. On real data its heading is no better than PCA; it earns its place by catching "
        "fragments and merges. Length: 90th percentile of the track's gate-passed fits + 0.12 m (a quantile, not the "
        "maximum, so one bad fit can't set it); 4.14 m median prior if fewer than 5 frames pass."))

    # ---------------- III. EXPERIMENTAL EVALUATION (friend: sub-dividers 1..5; results 3.1 -> 3.1.x, 3.2 -> 3.2.x)
    D("III. Experimental Evaluation", "3.1 Dataset and Protocol")
    C(III, "3.1.1 Datasets",
      "Real SemanticKITTI scans for detection; synthetic ShapeNet cars for completion.",
      lay_datasets, None,
      S("tab:dataset-roles, tab:sensor-spec, tab:dataset, verbatim. Seq 08 is the classifier's validation "
        "sequence, not a test set; LOSO (3.2.3) is the leakage control."))
    C(III, "3.1.2 Target Eligibility and Ground Truth",
      "A car is scored only if enough of its points survive preprocessing.",
      lay_eligibility, None,
      S("ch4 §eligibility — ≥10 surviving points; 25.5% of annotated cars excluded; "
        "recall vs ALL annotated cars ≈ 0.55 (ch4, ch5) vs 0.732 eligible"))
    SUB("3.2 Experimental Setup")
    C(III, "3.2.1 Baselines",
      "What each part of the pipeline is compared against.",
      lay_baselines, None,
      S("ch4 §Baseline Methods; tab:positioning (ch2). Why no head-to-head with supervised networks: they "
        "report all-class mIoU / panoptic quality and train on dense per-point labels; we report a car-only "
        "point-level F1. Putting them in one column would compare numbers the metrics do not support; the "
        "accuracy gap is real and stated in the thesis text. Completion has no external baseline on real "
        "data, so it is compared with its own input (paired), a dense reference (upper bound) and PoinTr."))
    C(III, "3.2.2 Metrics",
      "Detections are matched to real cars by point overlap (IoU); completions are scored by an independent car classifier.",
      lay_metrics, None,
      T("0:30", "fig:bg-iou; ch4 §Evaluation Metrics; ch4 §plausibility. Say: a detection counts as a car when it shares at least 30% of its points with a real car (intersection over union); best overlap first, one-to-one. Matched is a true positive, an extra detection a false positive, a missed car a false negative, counted over every frame. For completion there is no ground truth, so an independent judge decides whether each output looks like a car (next section). Don't explain why 0.3. HOOK H2 → slide 46; answer in qa_bank.md."))
    C(III, "3.2.3 Data Split: Leave-One-Sequence-Out",
      "Train on ten sequences, test on the eleventh, and repeat for every sequence.",
      lay_loso, None,
      S("ch4 §Dataset Splits, §LOSO; tab:loso, tab:headline (scripts/make_loso_grid.py). Fold-08 and the final classifier share data, architecture and the 15-epoch regime; they differ only in which epoch is kept. Pooled sums the counts (0.717); the macro mean of the 11 folds is 0.632 ± 0.119, pulled down by the small, sparse seqs 01 and 04. Seq 00 fold: P 0.921 / R 0.758 / F1 0.831."))
    SUB("3.3 Detection Results")
    C(III, "3.3.1 Overall Detection Results", "F1 0.812 on seq 08 (0.793 cross-validated), and each stage has a measurable job.",
      lay_results, None,
      T("1:00", "D9 = tab:ablation-stage, numbers verbatim (scripts/make_d9_funnel.py); table = tab:headline + "
               "tab:loso pooled row. Story, left to right: the geometric filter keeps most cars (recall 0.773) but 4 "
               "of 5 detections are false positives (precision 0.199); the classifier removes about 104,000 false "
               "positives (precision 0.846) and pays some recall (0.676); the track filter wins back 1,968 true "
               "positives and drops more false positives (0.911 / 0.732 / F1 0.812). Counts are per-frame matches "
               "over the sequence, not distinct cars. Table row 2 = tab:loso pooled (each fold's classifier never "
               "saw its test sequence). If asked whether we selected and report on seq 08: the seq-08 fold of "
               "LOSO (classifier blind to seq 08) gets 0.888 / 0.736 / 0.805, so selection bought about 2 "
               "precision points and no recall (tab:headline)."))
    C(III, "3.3.2 Detections on a Real Frame",
      "One seq-08 frame: found cars, false detections and missed cars.",
      lay_image, "thesis_bev_detections.png",
      S("fig:bev-detections (thesis seq08_bev_detections.png, unchanged). False positives about 0.6 per frame (2,513 / 4,071)."))
    C(III, "3.3.3 What Limits Recall",
      "Clustering splits cars into fragments; the classifier is not the limit.",
      lay_recall_limit, None,
      T("1:00", "D10 = fig:failure-zooms frame-250 row (scripts/make_d10_crop.py); split rate by range = ch4 §split analysis (0.10 m grid); recall by range = tab:distance. Say: the geometric stage alone already finds 77% of cars, more than the full pipeline, so the classifier is not what limits recall. The limit is clustering: a third of cars are split into pieces (33.7%, tab:splitrate). Here a parked car is broken up; one thin slice is detected but overlaps the car too little, so this one car costs a false positive and a miss. Splits are worst close by: half of the cars within 10 m, so near-range recall is lower than at 10 to 20 m. We tried six ways to repair the splitting; only one worked: a coarser clustering grid, recall 0.710 to 0.732. (Don't list the others.) HOOK H3 → backup 56; answer in qa_bank.md."))
    C(III, "3.3.4 Accuracy vs Range", "Recall peaks at 10–20 m: near cars get split, far cars are too sparse.",
      lay_distance, None, S("tab:distance, verbatim; split rates from ch4 §split analysis. Stride-20 sample, per frame, track filter off (striding breaks the tracker), so pooled recall 0.671 is not the 0.732 headline. Mean IoU of matched cars rises with range: the far cars found are compact single clusters."))
    C(III, "3.3.5 Detection Result Analysis", "What each detection stage contributes, and where it fails.",
      lay_cards_text, [("Clustering", "The recall limit: a third of cars are split into pieces."),
                       ("Geometric filter", "Keeps most cars (recall 0.773), but 4 in 5 detections are junk."),
                       ("Classifier", "The precision step: removes about 104,000 false positives."),
                       ("Track filter", "Wins back 1,968 true positives and drops 1,790 false positives.")],
      S("ch4 §Result Analysis, tab:ablation-stage, tab:splitrate (33.7% split at 0.10 m). Neither the classifier nor the tracker can create a detection, so recall can never exceed the 0.773 of the geometric stage. Precision-saturated (about 0.6 FP per frame), recall-limited."))
    SUB("3.4 Completion Results")
    C(III, "3.4.1 Completion Example: One Real Car",
      "A single-frame scan sees one corner of the car; completion fills in the hidden sides.",
      lay_image, "d13_completion_example.png",
      T("0:30", "D13 (scripts/pick_completion_examples.py, render_completion_candidates.py); "
                "concept: fig:completion-shape. Picked by a fixed rule, not by eye: among the cars the "
                "judge rejects as a raw scan and accepts once completed, with a typical number of input "
                "points (middle half of all 2,822), 12 drawn at random; this is one of them. Rotated so "
                "the length runs left-right, for display only. Say: blue is what the sensor saw in one "
                "frame, orange is what the network adds. Next slide: how we know the orange is a plausible car "
                "without ground truth."))
    C(III, "3.4.2 Judging Completion Without Ground Truth",
      "With no true shape on real cars, an independent judge decides what looks like a car.",
      lay_judge, None,
      T("1:00", "D12a (user, draw.io; scripts/make_d12a_scaffold.py) + D12b = tab:plausibility, verbatim; "
                "ch4 §Real Data Plausibility; protocol of Chen et al. 2020. Say: on a real car there is no "
                "true complete shape to measure against, and measuring distance to the scan itself rewards "
                "doing nothing (the raw scan scores a perfect zero, slide 39). So, as Chen et al. do, we ask "
                "an independent judge: a separate program trained on a different collection of 3D models. "
                "(Don't list why it's fair.) Every cloud gets the same preparation and the judge says car or not. Raw single-frame "
                "scans: judged a car 25% of the time. Completed: 79%. Dense multi-frame scans of parked cars, "
                "the best real shapes we have: 81%. Completion brings single-frame cars almost to the level "
                "of the best real shapes. Paired on the same cars: 1,615 switch to accepted, 115 the other way. "
                "HOOK H4 → backup 58 (and 59 if asked why distance rewards doing nothing); answer in qa_bank.md. "
                "Judge test accuracy 0.9214 (pass mark 0.88): keep for the answer, not the talk."))
    C(III, "3.4.3 Synthetic Fidelity",
      "On held-out synthetic cars, the network reproduces the complete shape.",
      lay_synthetic, None,
      S("ch4 §Synthetic Fidelity; D16 = scripts/make_d16_synthetic_example.py: a car from PCN's validation "
        "split (never trained on), picked by a fixed rule (the median-Chamfer car of the evaluated set), "
        "not by eye; its own Chamfer 0.16 m / F 0.77 is that one car. Chamfer distance: average "
        "nearest-point distance in both directions, added (metres). F-score @ 0.1 m: harmonic mean of the "
        "share of output points within 0.1 m of the true model and the share of the true model within "
        "0.1 m of the output. Best validation loss 0.1246 (coarse + 0.5 × fine Chamfer, normalised units; "
        "not the same number as 0.16 m)."))
    C(III, "3.4.4 Plausibility Robustness",
      "The plausibility gain holds on every sequence, for moving cars, and under a heading flip.",
      lay_plaus_robust, None,
      S("tab:plausibility-seq, verbatim; ch4 §plausibility (McNemar 1,615 vs 115; Wilcoxon on car "
        "probability p ≈ 0; movers 0.837 / 0.801, N = 602; heading flip Δ 0.009; parked subset N = 2,220: "
        "completed acceptance 0.770, car probability 0.730 vs ceiling 0.728, about 96% of the raw-to-ceiling "
        "acceptance gap). Moving vs parked: by net centroid displacement of the track. PCN never saw any "
        "real sequence, so consistency across all eleven is a generalisation result. Why 'not just "
        "density': the dense reference is dense AND correctly shaped, so matching it means the completed "
        "shapes are car-like, not merely dense."))
    C(III, "3.4.5 Completion Result Analysis",
      "What the completion evidence shows, and what it does not claim.",
      lay_completion_analysis, None,
      S("tab:fidelity, verbatim (UHD p95 column left out: 0.000 / 0.000 / 0.183); ch4 §Reference-Free "
        "Fidelity, §Scope of the Completion Claim. UCD: for each scan point, the distance to the nearest "
        "point of the cloud, averaged; UHD: the largest such distance. Direction scan → cloud, per-car "
        "medians, as Chen et al. and P2C apply them. Mirrored = the scan plus its mirror copy: it contains "
        "the scan, so it also scores 0. This is why the unpaired completion literature pairs these metrics "
        "with a plausibility term."))
    SUB("3.5 Runtime Evaluation")
    C(III, "3.5.1 Runtime by Stage",
      "The pipeline runs offline; the classical CPU stages dominate the time per scan.",
      lay_runtime, None, S("tab:runtime, verbatim (RTX 3070 Ti / Ryzen 7 7800X3D). Classifier 33.7 ms for about 45 clusters, about 0.75 ms each. Completion adds about 13 ms per completed car, outside this budget."))
    C(III, "3.5.2 Runtime Optimisation",
      "Fewer RANSAC iterations and a coarser clustering voxel cut the time per scan by about a third.",
      lay_runtime_opt, None, S("ch3 §RANSAC (1,000 → 300 iterations); ch4 §runtime evaluation (921.3 → 623.5 ms full seq; 650.6 ms on the stride-20 sample). The post-ground cloud is already sparse, so a coarser grid compresses it little; Python HDBSCAN keeps a floor of about 185 ms."))
    SUB("3.6 Ablation Study")
    C(III, "3.6.1 Stage Ablation", "Adding one filtering stage at a time shows what each contributes.",
      lay_stage_ablation, None,
      S("tab:ablation-stage, verbatim. Same numbers as the slide 29 chart, plus mean IoU and the raw counts. "
        "Counts are per-frame matches over the full sequence, not distinct cars."))
    C(III, "3.6.2 Clustering Alternatives and Resolution",
      "No faster clusterer matches HDBSCAN's recall; a coarser resolution reduces fragmentation.",
      lay_clustering_tables, None,
      S("tab:clustering, tab:splitrate. Mean IoU column of tab:clustering left out. "
        "If asked 'why not an even coarser grid, e.g. 0.20 m?': merges grow as the grid coarsens "
        "(56 → 284), so we expect more fused cars; 0.20 m was NOT tested, say 'expect', not 'found'. "
        "If asked 'did you tune DBSCAN's radius?': yes, an epsilon sweep; the best fixed radius "
        "(Euclidean, 0.4 m) reaches F1 0.800 on seq 00, still below HDBSCAN's 0.835 (ch4 §clustering)."))
    C(III, "3.6.3 IoU-Threshold Sensitivity", "F1 moves by under 3 points between IoU 0.25 and 0.50.",
      lay_iou_sensitivity, None, S("tab:iou-sensitivity, verbatim. 0.5 is the usual box threshold; ours is point-level. Pre-registered: a threshold-sensitive result would have been reported as a finding, not tuned away."))
    C(III, "3.6.4 Synthetic Pretraining",
      "Synthetic pretraining adds nothing once enough real clusters are mined.",
      lay_pretraining, None,
      S("ch4 sec:ablation-stagea (macro-F1 0.9285 vs 0.9225; pipeline F1 identical to three decimals); "
        "tab:cross-domain. Don't mix the metrics: 0.9285 is macro-F1 (mean of car and not-car F1), 0.885 is "
        "car-class F1. Why keep it in the thesis: it changes the recipe, so it is reported as a negative result."))
    C(III, "3.6.5 Completion Ablations",
      "Input quality and normalisation matter more than the choice of network.",
      lay_completion_ablations, None,
      S("ch4 §Completion Ablation Studies (sec:abl-completion), one row each. Normalisation: identical "
        "checkpoint, only the input frame differs, so the blob-vs-car gap is a normalisation effect. "
        "Fine-tuning: the early blobs were blamed on a domain gap, so PCN was fine-tuned on LiDAR-like "
        "simulated scans (noise augmentation, a virtual Velodyne; findings #16, #17, #19): still blobs, "
        "because the cause was the input frame. Real-data fine-tuning was never run: it is future work (slide 52); do not say it was tried. Gate: its cost is some usable cars "
        "rejected. Length: a residual far-end shortfall on the longest cars remains (ch5). PoinTr: same "
        "held-out synthetic cars; the remaining real-data gap is centre estimation, which no network "
        "fixes."))

    # ---------------- IV. CONCLUSION & FUTURE WORK (friend: 1 summary, 2 limitations, 3 discussion & future)
    D("IV. Conclusion & Future Work")
    C(IV, "4.1 Research Summary",
      "Four findings from a modular framework that needs zero additional manual annotation.",
      lay_cards_text, [("Detection on real LiDAR", "F1 0.812 on seq 08; 0.793 cross-validated over all 11 sequences"),
                       ("Plausible completions", "Independent judge's car acceptance: 0.253 raw → 0.785 completed"),
                       ("Clustering-bound recall", "HDBSCAN splits 33.7% of cars; the classifier is not the limit"),
                       ("Redundant synthetic pretraining", "Enough auto-mined real clusters make it unnecessary")],
      T("0:45", "ch5 §principal outcomes, §diagnostic findings, Final Remarks"))
    C(IV, "4.2 Limitations", "Four limits on what the results show.",
      lay_cards_text, [("Dense scenes only", "Highway seq 01 recall 0.336 vs 0.758 on seq 00 (precision still 0.951)"),
                       ("Recall ceiling", "0.773 before any learned stage"),
                       ("Plausibility only", "Fidelity metrics do not corroborate the completion gain"),
                       ("Offline", "623.5 ms per frame; real time is out of scope")],
      T("0:30", "ch5 §limitations (disc-architectural, disc-completion-scope, disc-eval-validity). Say the four "
               "titles with their numbers; on Offline say only '623.5 ms per frame, offline by design' and move on. "
               "HOOK H5 → slides 41–42; answer in qa_bank.md."))
    C(IV, "4.3 Discussion & Future Work", "Three next steps, each scoped by a finding.",
      lay_cards_text, [("Real-time variant", "A native or GPU HDBSCAN implementation"),
                       ("Real-data fine-tuning", "Close the synthetic-to-real gap for completion"),
                       ("Downstream detection", "Test whether completed cars improve 3D detection")],
      T("0:30", "ch5 §future work"))
    end(prs, nxt())

    # ---------------- Backups (after THE END, like the friend's appendix)
    C(BK, "LiDAR Point Cloud Properties",
      "A LiDAR scan is an unordered set of 3D points that gets sparser with range and sees each object from one side.",
      lay_cards_text, [("Unordered", "No fixed order, unlike the pixels of an image"),
                       ("Sparser with range", "Far objects get only a handful of points"),
                       ("One-sided", "Each object is seen from one side only")],
      S("ch2 §LiDAR point cloud; tab:sensor-spec; S1 notes A1"))
    C(BK, "How HDBSCAN Works",
      "It tries every radius at once, and keeps a group whole if it holds together longer than its pieces would.",
      lay_image, "hdbscan_explainer.png",
      S("backup for 'how does HDBSCAN work?'. TOY 2D example (scripts/make_hdbscan_explainer.py), "
        "real hdbscan library with our settings (min_samples 5, min_cluster_size 10); NOT a thesis result. "
        "Say: (1) each point measures how crowded its neighbourhood is: the distance to its 5th nearest "
        "point (itself included); (2) HDBSCAN then looks at every radius at once: at a small radius only "
        "the dense near car forms, at a medium one both cars form, at a large one they fuse; (3) shrinking "
        "the radius, it keeps a group whole if it holds together longer than its pieces would. Each car is "
        "judged only against its own pieces, never against the other car, so the dense near car and the "
        "sparse far car both survive; a fixed-radius DBSCAN would have to pick one column. "
        "Toy numbers: near car 335.7 vs its pieces 19.2 -> whole; far car 24.1, no pieces -> kept "
        "(S2 notes A4). "
        "Precision if pressed: the 'radius' is the mutual reachability distance max(core a, core b, "
        "distance a-b), and 'longest' is stability, lifetime measured in 1/radius and summed over points; "
        "groups under 10 points are noise. ch2 §density-based clustering; S2 notes A4."))
    C(BK, "Repair Strategies and Why They Fail",
      "Each repair either over-fits seq 00 or trades precision for recall.",
      lay_strategies, None,
      S("backup for 'did you try to fix the splitting?'. tab:strategies, verbatim; ch4 §Limits of the Repair "
        "Strategies. Pre-promotion seq-00 baseline (0.05 m grid), so the baseline F1 0.844 is not the final "
        "pipeline. Shared limit: density clustering has no notion of objectness, so no proximity or size rule "
        "can tell 'rejoin two pieces of one car' from 'keep two adjacent cars apart'. Resolution worked because "
        "0.10 m closes gaps inside a car before it closes gaps between cars."))
    C(BK, "Normalisation Ablation on a Real Car",
      "Same network, same scan: only the frame the scan is put in differs.",
      lay_image, "b57_normalisation_ablation.png",
      S("backup for 'why the four geometry steps?' or 'why not fine-tune on real data?'. ch4 §Completion "
        "Ablation Studies, Inference normalisation + Sim-to-real fine-tuning. Figure: "
        "scripts/make_b57_normalisation_ablation.py, car C (fixed in advance), checkpoint pcn_kitti_best on both "
        "sides; left = the original path verbatim (scan's own centroid and radius, 3D PCA, longest axis as X), "
        "right = the production output. Left output spans about 5.1 × 3.5 × 2.6 m, right 3.9 × 1.8 × 1.7 m "
        "(illustration, not a thesis number). Say: the network was never the problem; the input frame was. "
        "Early blobs were first blamed on the synthetic-to-real gap, but retraining on more realistic "
        "simulated scans did not help; putting the scan in the training frame did."))
    C(BK, "Is the Judge Independent?",
      "Each way the judge could favour us, and what rules it out.",
      lay_judge_independent, None,
      S("backup for any challenge to slide 36. ch4 §Real Data Plausibility; Chen et al. 2020 use a "
        "ModelNet40 judge for the same reason. Pre-registration and the 0.88 gate: written before scoring. "
        "ModelNet40 has no truck, bus or van class, so 'car' is not split with near-vehicle classes."))
    C(BK, "Fidelity Metrics: UCD and UHD",
      "They compare the output with its own scan, so a copy of the scan scores best.",
      lay_ucd_uhd, None,
      S("backup for 'why not the standard metrics?'. ch2 sec:th-distances, fig:bg-fidelity (redrawn as a toy, "
        "scripts/make_b59_ucd_uhd.py, numbers are the toy's own); tab:fidelity. Used on KITTI cars by Chen et "
        "al. 2020 (UHD) and P2C 2023 (UCD). Reported in the thesis, with the caveat that they reward "
        "under-completion."))
    C(BK, "Worked Example: IoU Matching",
      "A detection counts as a car when it shares at least 30% of the points.",
      lay_iou_example, None,
      S("backup for 'how exactly is a detection matched?'. ch2 eq. IoU; ch4 §Evaluation Metrics; fig:bg-iou "
        "(slide 26). Worked example = illustration (S4 notes A4). A split car costs twice: a missed car and "
        "often a false detection. Counts are per frame, micro-averaged."))
    C(BK, "Pipeline Configuration",
      "Every number in the thesis comes from these frozen values.",
      lay_pipeline_config, None,
      S("backup for 'what parameters?'. tab:pipeline-config, verbatim . Fixed before "
        "evaluation; nothing tuned on the reported results. Completion constants: slide 20 "
        "(tab:completion-config)."))

    add_vi_script(prs)

    out = OUT
    try:
        prs.save(out)
    except PermissionError:  # deck open in PowerPoint locks the file
        out = OUT.with_name(OUT.stem + "_new.pptx")
        prs.save(out)
        print(f"NOTE: {OUT.name} is open/locked; saved to {out.name} instead")
    print(f"wrote {out} ({len(prs.slides)} slides)")
    for i, sl in enumerate(prs.slides, 1):
        txt = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        print(i, (txt[0] if txt else "").replace("\n", " / ")[:70])


if __name__ == "__main__":
    build()
