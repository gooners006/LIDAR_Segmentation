"""Print the deck as a markdown slide table (# | title | speaker notes) for docs/defense/outline.md.

Usage: .venv/Scripts/python.exe docs/defense/scripts/deck_table.py docs/defense/deck_v2.pptx > table.md
"""
import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding="utf-8")
p = Presentation(sys.argv[1])
print("| # | Slide | Talk path / source (from speaker notes) |")
print("|---|---|---|")
for i, sl in enumerate(p.slides, 1):
    txt = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    title = txt[0].replace("\n", " · ") if txt else ""
    if i == 1:
        title = "Cover"
    note = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    if note.startswith("divider"):
        title = "▌" + title
    print(f"| {i} | {title} | {note} |")
