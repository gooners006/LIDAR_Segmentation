"""Dump every slide's text, table cells, image names and speaker notes (used for wording passes).

Usage: PYTHONIOENCODING=utf-8 .venv/Scripts/python.exe docs/defense/scripts/dump_deck.py docs/defense/deck_v2.pptx
"""
import sys
from pptx import Presentation
p = Presentation(sys.argv[1])
for i, s in enumerate(p.slides, 1):
    print(f"=== {i}")
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print("  T:", sh.text_frame.text.replace("\n", " | "))
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                print("  R:", " | ".join(c.text for c in r.cells))
        if sh.shape_type == 13:
            print("  IMG:", sh.name)
    if s.has_notes_slide:
        n = s.notes_slide.notes_text_frame.text.strip()
        if n: print("  N:", n.replace("\n", " / "))
