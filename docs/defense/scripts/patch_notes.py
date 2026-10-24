"""Replace a phrase in one slide's speaker notes in place, run by run (keeps the rest of the deck untouched).

Usage: PYTHONIOENCODING=utf-8 .venv/Scripts/python.exe docs/defense/scripts/patch_notes.py <deck.pptx> <slide#> <old> <new> [<old> <new> ...]
Fails if any phrase isn't found inside a single run exactly once.
"""
import sys
from pptx import Presentation

deck, n, pairs = sys.argv[1], int(sys.argv[2]), sys.argv[3:]
prs = Presentation(deck)
runs = [r for p in prs.slides[n - 1].notes_slide.notes_text_frame.paragraphs for r in p.runs]
for old, new in zip(pairs[::2], pairs[1::2]):
    hits = [r for r in runs if old in r.text]
    assert len(hits) == 1 and hits[0].text.count(old) == 1, f"slide {n}: {old[:50]!r} found {len(hits)}x"
    hits[0].text = hits[0].text.replace(old, new)
prs.save(deck)
print(f"{deck}: slide {n} patched ({len(pairs) // 2} phrase(s))")
