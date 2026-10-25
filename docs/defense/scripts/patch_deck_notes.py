"""Patch speaker notes in place in the built decks, without a full rebuild (keeps any manual
slide edits). Talk slides become: VI script (+ VI extra), EN script (+ EN extra), then the English
planning line; support / backup / divider slides become: VI note (notes_vi.py), then the English
line. Both are composed by build_deck's own functions, so a rebuild gives the same notes.
Idempotent: the English line is recovered from whichever layout the notes are in.

Usage: .venv/Scripts/python.exe docs/defense/scripts/patch_deck_notes.py
"""
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).parent))
from build_deck import (SUPPORT_VI, compose_notes, compose_support_notes,  # noqa: E402
                        read_talk_script, slide_title)

ROOT = Path(__file__).resolve().parents[1]
DECKS = [ROOT / "deck_v2.pptx", ROOT / "deck_v2_embedded.pptx"]

# English-line corrections already made in build_deck.py, applied to the built decks.
EN_FIXES = {
    46: ("0.5 is the usual box threshold; ours is point-level.",
         "0.3 was the evaluator default from the start, not tuned on results; the point is that the "
         "conclusion holds at 0.25 and 0.5 (H2, qa_bank.md). Don't say 'with boxes 50% is the rule': "
         "KITTI's car box benchmark uses 0.7."),
}


def english_line(text, starts):
    """The one paragraph that starts with any of `starts`, wherever it sits in the notes."""
    hits = [p for p in text.split("\n\n") if p.startswith(starts)]
    assert len(hits) == 1, f"expected one English line, got {len(hits)}"
    return hits[0]


blocks = read_talk_script()
for deck in DECKS:
    prs = Presentation(deck)
    for n, b in blocks.items():
        tf = prs.slides[n - 1].notes_slide.notes_text_frame
        tf.text = compose_notes(n, english_line(tf.text, "★"), b)
    for n, (want, _) in SUPPORT_VI.items():
        sl = prs.slides[n - 1]
        assert slide_title(sl).splitlines()[0].upper() == want.upper(), f"{deck.name} slide {n}: title changed"
        tf = sl.notes_slide.notes_text_frame
        meta = english_line(tf.text, ("○", "divider:"))
        if n in EN_FIXES:
            meta = meta.replace(*EN_FIXES[n])
        tf.text = compose_support_notes(n, meta)
    prs.save(deck)
    print("patched", deck.name, f"({len(blocks)} talk + {len(SUPPORT_VI)} other slides)")
